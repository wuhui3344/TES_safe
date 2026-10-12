"""报告与追溯（需求模块 4）。

导出 CSV / Excel / PDF，记录输入参数、计算公式、规则来源、标准版本、风险项与待验证项目。
所有导出函数均为纯函数（输入数据 → bytes/str），便于单元测试与复用。

依赖策略（刻意为之）：

* CSV —— 标准库 ``csv``；
* Excel —— 仅用 ``openpyxl``，**不用 pandas**（pandas+numpy 会让打包体积增加约 80 MB）；
* PDF —— 本模块输出自包含 HTML（纯标准库），由界面端（Qt ``QTextDocument``）渲染成 PDF，
  从而无需 ``reportlab``。
"""

from __future__ import annotations

import csv
import html
import io
from datetime import datetime
from typing import Dict, List, Optional, Sequence

from .models import ProjectConstraints, SampleAnalysis, StimulationParameters
from .safety_rules import (
    STATUS_INSUFFICIENT,
    STATUS_PASS,
    STATUS_VIOLATION,
    RuleCheck,
    rule_catalog_records,
)
from .waveforms import CalculationResult


def build_parameter_record(params: StimulationParameters, constraints: ProjectConstraints) -> Dict[str, object]:
    record: Dict[str, object] = {
        "记录时间": datetime.now().isoformat(timespec="seconds"),
        "刺激类型": params.mode,
        "波形定义": params.waveform,
        "电流幅值定义": params.current_definition,
        "电流输入值_mA": params.current_value_mA,
        "频率_Hz": params.frequency_hz,
        "单相脉宽_us": params.pulse_width_us,
        "每秒脉冲数_pps": params.pulses_per_second,
        "相数": params.phases,
        "电极形式": params.electrode_type,
        "有效接触面积_cm2": params.area_cm2,
        "通道数": params.channels,
        "单次时长_min": params.duration_min,
        "每天会话数": params.sessions_per_day,
        "负载阻抗_ohm": params.impedance_ohm,
        "合规电压_V": params.compliance_voltage_V,
        "项目电流幅值约束_启用": constraints.use_current_amplitude,
        "项目电流幅值上限_mA": constraints.current_amplitude_limit_mA,
        "项目电流密度约束_启用": constraints.use_current_density,
        "项目电流密度上限_mA_cm2": constraints.current_density_limit_mA_cm2,
        "项目电荷密度约束_启用": constraints.use_charge_density,
        "项目电荷密度上限_mC_cm2": constraints.charge_density_limit_mC_cm2,
        "项目时长约束_启用": constraints.use_duration,
        "项目时长上限_min": constraints.duration_limit_min,
    }
    return record


def build_calculation_record(result: CalculationResult) -> Dict[str, object]:
    return {
        "波形族": result.family_label,
        "峰值电流_mA": _num(result.peak_current_mA),
        "RMS电流_mA": _num(result.rms_current_mA),
        "峰峰值电流_mA": _num(result.pp_current_mA),
        "峰值电流密度_mA_cm2": _num(result.current_density_peak_mA_cm2),
        "RMS电流密度_mA_cm2": _num(result.current_density_rms_mA_cm2),
        "峰值电流密度_A_m2": _num(result.current_density_peak_A_m2),
        "净电荷_mC": _num(result.net_charge_mC),
        "绝对电荷_mC": _num(result.absolute_charge_mC),
        "每相电荷_mC": _num(result.charge_per_phase_mC),
        "净电荷密度_mC_cm2": _num(result.charge_density_net_mC_cm2),
        "绝对电荷密度_mC_cm2": _num(result.charge_density_absolute_mC_cm2),
        "估算峰值电压_V": _num(result.peak_voltage_V),
        "估算负载功率_W": _num(result.power_W),
        "估算能量_J": _num(result.energy_J),
        "每日总时长_min": _num(result.daily_total_duration_min),
        "归一化比较电流_mA": _num(result.reference_current_mA),
        "归一化基准": result.reference_current_basis,
        "计算说明": "；".join(result.notes),
    }


def build_sample_record(sample: Optional[SampleAnalysis]) -> Dict[str, object]:
    if sample is None:
        return {"实测波形": "未导入"}
    data = sample.to_dict()
    data["实测波形"] = "已导入"
    return data


def build_check_records(checks: List[RuleCheck]) -> List[Dict[str, object]]:
    return [c.to_dict() for c in checks]


def build_summary(checks: List[RuleCheck]) -> Dict[str, int]:
    summary = {STATUS_PASS: 0, STATUS_VIOLATION: 0, STATUS_INSUFFICIENT: 0}
    for c in checks:
        summary[c.status] = summary.get(c.status, 0) + 1
    summary["合计"] = len(checks)
    return summary


def _num(value: Optional[float]):
    if value is None:
        return ""
    return round(value, 6)


def rows_to_csv_bytes(rows: List[Dict[str, object]]) -> bytes:
    """把多行字典导出为 CSV。

    字段名取**所有行的并集**（按首次出现顺序），因此允许各行拥有不同字段
    （例如“输入参数”行与“计算结果”行）；缺失字段留空。
    """
    if not rows:
        return b""
    fieldnames: List[str] = []
    for row in rows:
        for key in row:
            if key not in fieldnames:
                fieldnames.append(key)
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=fieldnames, restval="")
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8-sig")


def merge_records_for_csv(sections: Dict[str, Dict[str, object]]) -> List[Dict[str, object]]:
    """把多个（节名 → 记录）字典展平为“分类 / 项目 / 数值”行，便于阅读的 CSV 导出。"""
    rows: List[Dict[str, object]] = []
    for section, record in sections.items():
        for key, value in record.items():
            rows.append({"分类": section, "项目": key, "数值": value})
    return rows


def _transpose(record: Dict[str, object]) -> List[Dict[str, object]]:
    return [{"项目": k, "数值": v} for k, v in record.items()]


def build_excel_sheets(
    params: StimulationParameters,
    constraints: ProjectConstraints,
    result: CalculationResult,
    checks: List[RuleCheck],
    sample: Optional[SampleAnalysis] = None,
) -> Dict[str, List[Dict[str, object]]]:
    """构造 Excel 各工作表数据（与导出引擎无关，便于测试）。"""
    sheets: Dict[str, List[Dict[str, object]]] = {
        "输入参数": _transpose(build_parameter_record(params, constraints)),
        "计算结果": _transpose(build_calculation_record(result)),
        "规则筛查": build_check_records(checks),
        "规则依据": rule_catalog_records(),
    }
    if sample is not None:
        sheets["实测波形"] = _transpose(build_sample_record(sample))
    return sheets


def _display_width(text: str) -> int:
    """粗略估算显示宽度：中日韩字符按 2 计，用于 Excel 列宽。"""
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in text)


def sheets_to_excel_bytes(sheets: Dict[str, List[Dict[str, object]]]) -> Optional[bytes]:
    """用 openpyxl 生成多工作表 Excel；未安装 openpyxl 时返回 None。

    刻意不使用 pandas：本工具只需要写出几个二维表，而 pandas+numpy
    会让打包产物增加约 80 MB，并明显拖慢启动。
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:  # pragma: no cover - 可选依赖
        return None

    try:
        workbook = Workbook()
        workbook.remove(workbook.active)
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="4472C4")

        for name, rows in sheets.items():
            worksheet = workbook.create_sheet(title=str(name)[:31])
            if not rows:
                continue
            headers = list(rows[0].keys())
            worksheet.append(headers)
            for cell in worksheet[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(vertical="center")
            for row in rows:
                worksheet.append([row.get(h, "") for h in headers])
            for index, header in enumerate(headers, start=1):
                width = max(
                    [_display_width(str(header))]
                    + [_display_width(str(row.get(header, ""))) for row in rows]
                )
                worksheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 10), 60)
            worksheet.freeze_panes = "A2"

        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()
    except Exception:  # pragma: no cover
        return None


# --------------------------------------------------------------------------- #
# HTML 报告（供界面端渲染为 PDF，或直接另存为 .html）
# --------------------------------------------------------------------------- #

_STATUS_COLORS = {
    STATUS_VIOLATION: "#c0392b",
    STATUS_INSUFFICIENT: "#b9770e",
    STATUS_PASS: "#1e8449",
}


def _html_table(
    headers: Sequence[str],
    rows: Sequence[Sequence[object]],
    widths: Optional[Sequence[int]] = None,
    status_column: Optional[int] = None,
) -> str:
    """生成表格。

    刻意使用 width / bgcolor / font 等“老式”属性而不是 CSS：
    Qt 的 ``QTextDocument`` 富文本引擎只支持 CSS 的一个子集（如不支持
    ``border-collapse``），而属性写法在 Qt 与浏览器里都能正确渲染。
    """
    count = max(len(headers), 1)
    if widths is None:
        widths = [max(8, 100 // count)] * count

    head = "".join(
        f'<th width="{widths[i]}%" bgcolor="#4472C4">'
        f'<font color="#ffffff"><b>{html.escape(str(h))}</b></font></th>'
        for i, h in enumerate(headers)
    )

    body_parts: List[str] = []
    for row in rows:
        cells: List[str] = []
        for index, value in enumerate(row):
            text = html.escape(str(value))
            if index == status_column and str(value) in _STATUS_COLORS:
                text = f'<font color="{_STATUS_COLORS[str(value)]}"><b>{text}</b></font>'
            cells.append(f"<td>{text}</td>")
        body_parts.append("<tr>" + "".join(cells) + "</tr>")

    return (
        "<table width='100%' border='1' cellspacing='0' cellpadding='4' "
        f"bgcolor='#ffffff'><tr>{head}</tr>{''.join(body_parts)}</table>"
    )


def build_report_html(
    parameter_record: Dict[str, object],
    calculation_record: Dict[str, object],
    check_records: List[Dict[str, object]],
    sample_record: Optional[Dict[str, object]] = None,
    rule_records: Optional[List[Dict[str, object]]] = None,
) -> str:
    """生成自包含 HTML 审查报告（纯标准库，无 reportlab / pandas）。

    参数全部是“已序列化的普通字典”，因此可直接被缓存与单测，不依赖领域对象。
    """
    if rule_records is None:
        rule_records = rule_catalog_records()

    parts: List[str] = [
        "<html><head><meta charset='utf-8'></head><body style='font-family:sans-serif'>",
        "<h1 style='font-size:16pt'>经颅电刺激（tES）安全设计审查报告</h1>",
        "<p style='color:#c0392b'><b>重要声明：</b>本报告只计算物理量并筛查已录入的设计约束，"
        "<b>不输出“人体绝对安全”结论</b>。文献常见范围不是法规限值；未经风险分析、台架验证、"
        "必要的建模与临床/伦理审查，不得用于人体刺激或作为注册合规证明。</p>",
        "<h2 style='font-size:12pt'>一、输入参数</h2>",
        _html_table(["项目", "数值"], [[k, v] for k, v in parameter_record.items()], widths=[30, 70]),
        "<h2 style='font-size:12pt'>二、计算结果</h2>",
        _html_table(["项目", "数值"], [[k, v] for k, v in calculation_record.items()], widths=[30, 70]),
    ]

    if sample_record is not None:
        parts += [
            "<h2 style='font-size:12pt'>三、实测波形分析</h2>",
            _html_table(["项目", "数值"], [[k, v] for k, v in sample_record.items()], widths=[30, 70]),
        ]

    parts.append("<h2 style='font-size:12pt'>四、规则筛查</h2>")
    parts.append(_html_table(
        ["规则编号", "类别", "状态", "说明"],
        [[r.get("规则编号", ""), r.get("类别", ""), r.get("状态", ""), r.get("说明", "")]
         for r in check_records],
        widths=[14, 18, 16, 52],
        status_column=2,
    ))

    parts.append("<h2 style='font-size:12pt'>五、规则依据（可追溯性）</h2>")
    parts.append(_html_table(
        ["规则编号", "来源", "来源版本", "证据等级", "验证状态"],
        [[r.get("规则编号", ""), r.get("来源", ""), r.get("来源版本", ""),
          r.get("证据等级", ""), r.get("验证状态", "")] for r in rule_records],
        widths=[14, 30, 16, 20, 20],
    ))

    parts.append(
        "<p style='color:#555;font-size:9pt'><i>声明：任何项目阈值均应具有可追溯证据、"
        "批准记录与验证报告。本工具不能替代 IEC/GB/YY 标准全文、风险管理文件、"
        "独立安全测试、计算建模或临床/伦理审查。</i></p>"
    )
    parts.append("</body></html>")
    return "".join(parts)
